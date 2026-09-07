"""F01/F02 regressions through real readers, adapter, routes and BackgroundTasks.

Only physical storage and database sessions are replaced; no provider required.
"""
import asyncio
import csv
from datetime import datetime, timezone
from io import BytesIO, StringIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

from fastapi import BackgroundTasks
from openpyxl import Workbook, load_workbook
import pandas as pd
import pytest

from app.api.v1.endpoints import analysis as api
from app.core import database
from app.models import Analysis, Dataset
from app.schemas.analysis import AnalysisCreate
from app.schemas.capability import CapabilityPreflightRequest, DatasetSemanticState
from app.services import capability_compiler as compiler
from app.services import capability_execution_service as boundary
from app.services import capability_input_service as integrity
from app.services import dataset_io_service as reader
from app.services.capability_registry import CORE_OUTPUT

ROLES = {"entity_id": "user_id", "conversion_flag": "converted", "cohort": "cohort", "acquisition_channel": "channel"}
COLUMNS = ["user_id", "converted", "cohort", "channel"]


def rows(values):
    return [[f"u{i}", v, "previous_month" if i < 2 else "recent_month", "channel-a"] for i, v in enumerate(values)]


def csv_bytes(values, *, quoted=False, duplicate=False):
    output = StringIO(newline="")
    writer = csv.writer(output, quoting=csv.QUOTE_ALL if quoted else csv.QUOTE_MINIMAL)
    writer.writerow(COLUMNS)
    data = rows(values)
    writer.writerows(data + data[:1] if duplicate else data)
    return output.getvalue().encode()


def xlsx_bytes(values):
    book = Workbook()
    sheet = book.active
    sheet.append(COLUMNS)
    for row in rows(values):
        sheet.append(row)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def candidate(derived=False):
    evidence = ["cohort_cvr_comparison", "channel_cvr_comparison", "mix_within_decomposition"]
    return CapabilityPreflightRequest.model_validate({
        "candidate": {
            "schema_version": "candidate-plan@2", "plan_id": "integrity-plan", "plan_version": 1,
            "user_question": "Compare conversion and channel change",
            "capability_ref": {"capability_id": "conversion_decline_diagnosis", "version": "1"},
            "business_goal_ids": ["conversion_decline"], "required_evidence_types": evidence,
            "field_bindings": [
                {"dataset_ref": {"dataset_id": "derived" if derived else "users", "version": None},
                 "role": role, "column": col,
                 "provenance": {"source_dataset_id": "users", "source_column": col} if derived else None}
                for role, col in ROLES.items()
            ],
            "population_spec": {"grain": "user_order_detail" if derived else "unique_user", "entity_role": "entity_id", "scope": "all_input_users_in_selected_cohorts", "conversion_policy": "binary_required_missing_invalid"},
            "comparison_spec": {"baseline": "previous_month", "current": "recent_month"},
            "expected_output_contract": CORE_OUTPUT.model_dump(),
        },
        "requirements": {"business_goal_ids": ["conversion_decline"], "required_evidence_types": evidence},
    })


def metadata(dataset_id, extension="csv"):
    return SimpleNamespace(
        id=dataset_id, user_id="owner", is_deleted=False,
        filename=f"{dataset_id}.{extension}", storage_path=f"memory/{dataset_id}.{extension}",
        file_size=2048, schema=[{"name": col, "dtype": "int64" if col == "converted" else "object"} for col in COLUMNS],
        derivation_type=None, derivation_plan=None, transform_chain=None,
        parent_dataset_id=None, source_dataset_ids=None, derivation_risk_summary=None,
        display_name="display only", description="description only", ui_label="UI only",
        updated_at=datetime.now(timezone.utc),
    )


def saved_join():
    return {
        "id": "join-1", "source_analysis_plan_id": "plan-0", "relationship_set_id": "rels-1",
        "base_dataset_id": "users", "included_dataset_ids": ["users", "orders"],
        "join_steps": [{"left_dataset_id": "users", "right_dataset_id": "orders", "left_field": "user_id", "right_field": "user_id", "join_type": "left", "relationship_id": "r1", "relationship_status": "confirmed", "expected_cardinality": "one_to_many"}],
        "confirmed_relationships": [{"id": "r1", "source_dataset_id": "users", "source_field": "user_id", "target_dataset_id": "orders", "target_field": "user_id", "status": "confirmed", "expected_cardinality": "one_to_many", "risk_level": "high"}],
        "selected_fields": {"users": COLUMNS, "orders": ["user_id"]},
        "requires_confirmation": True, "warnings": [],
    }


class MemoryDB:
    def __init__(self, datasets):
        self.datasets, self.analyses = datasets, {}

    async def execute(self, statement):
        entity = statement.column_descriptions[0]["entity"]
        params = statement.compile().params
        if entity is Dataset:
            assert "datasets.user_id" in str(statement) and "datasets.is_deleted" in str(statement)
            obj = self.datasets.get(params["id_1"])
            if obj is not None and (obj.user_id != params["user_id_1"] or obj.is_deleted):
                obj = None
        else:
            assert entity is Analysis
            obj = self.analyses.get(params["id_1"])
        return SimpleNamespace(scalar_one_or_none=lambda: obj)

    def add(self, obj):
        self.analyses[obj.id] = obj

    async def commit(self):
        pass

    async def refresh(self, obj):
        obj.created_at = datetime.now(timezone.utc)

    async def __aenter__(self):
        return self

    async def __aexit__(self, *args):
        pass


def environment(monkeypatch, *, derived=False, content=None, extension="csv"):
    users = metadata("users", extension)
    datasets = {"users": users}
    files = {users.storage_path: content if content is not None else csv_bytes([0, 1, 1, 0])}
    if derived:
        joined = metadata("derived")
        joined.derivation_type = "join"
        joined.parent_dataset_id = "users"
        joined.source_dataset_ids = ["users", "orders"]
        joined.derivation_plan = saved_join()
        joined.derivation_risk_summary = {"base_grain": "unique_user", "result_grain": "likely_detail", "warnings": []}
        datasets["derived"] = joined
        files[joined.storage_path] = csv_bytes([0, 1, 1, 0], duplicate=True)
    db = MemoryDB(datasets)
    monkeypatch.setattr(reader.storage, "read", AsyncMock(side_effect=lambda path: files[path]))
    monkeypatch.setattr(database, "AsyncSessionLocal", lambda: db)
    return db, files


async def preflight(db, *, derived=False):
    response = await api.capability_preflight(candidate(derived), db, SimpleNamespace(id="owner"))
    return response.data


async def create(db, *, derived=False):
    outcome = await preflight(db, derived=derived)
    assert outcome.executable, outcome
    tasks = BackgroundTasks()
    request = AnalysisCreate(
        dataset_id="derived" if derived else "users", analysis_type=boundary.ANALYSIS_TYPE,
        params={"request": candidate(derived).model_dump(mode="json"), "expected_execution_spec_id": outcome.execution_spec.execution_spec_id},
    )
    response = await api.create_analysis(request, tasks, db, SimpleNamespace(id="owner"))
    assert len(tasks.tasks) == len(db.analyses) == 1
    return response.data, tasks, outcome.execution_spec


def mutate(kind, db, files, monkeypatch):
    users, derived = db.datasets["users"], db.datasets.get("derived")
    primary = derived or users
    if kind == "source_content":
        files[primary.storage_path] = csv_bytes([1, 1, 1, 0])
    elif kind in {"source_schema", "bound_dtype"}:
        primary.schema[1]["dtype"] = "boolean"
    elif kind.startswith("base_schema_"):
        users.schema[int(kind[-1])]["dtype"] = "changed-semantic-type"
    elif kind == "base_frame":
        text = files[users.storage_path].decode().splitlines()
        files[users.storage_path] = ("\n".join([text[0], *text[:0:-1]]) + "\n").encode()
    elif kind == "join_plan":
        derived.derivation_plan["id"] = "join-new"
    elif kind == "transform_chain":
        derived.transform_chain = [{"operation": "filter", "predicate": "user_id != u0"}]
    elif kind == "source_refs":
        derived.source_dataset_ids = ["users", "different-right-source"]
    elif kind == "base_ref":
        derived.parent_dataset_id = "different-base"
    elif kind == "relationship":
        derived.derivation_plan["confirmed_relationships"][0]["expected_cardinality"] = "one_to_one"
    elif kind == "grain":
        derived.derivation_risk_summary["result_grain"] = "different_grain"
    elif kind == "normalization_version":
        monkeypatch.setattr(integrity, "NORMALIZATION_POLICY", integrity.NORMALIZATION_POLICY.model_copy(update={"version": "2"}))
    elif kind == "reader_version":
        monkeypatch.setattr(integrity, "READER_CONTRACT", "capability-logical-reader@2")
    elif kind == "bound_semantics":
        primary.schema[1]["semantic_type"] = "different_semantics"
    elif kind == "column_removed":
        files[primary.storage_path] = files[primary.storage_path].replace(b"converted", b"removed")
    else:
        raise AssertionError(kind)


SEMANTIC_CASES = [
    (False, "source_content"), (False, "source_schema"), (False, "bound_dtype"),
    (False, "bound_semantics"), (False, "column_removed"),
    *[(True, f"base_schema_{i}") for i in range(4)],
    (True, "join_plan"), (True, "transform_chain"), (True, "source_refs"),
    (True, "base_ref"), (True, "relationship"), (True, "grain"), (True, "base_frame"),
    (True, "normalization_version"), (False, "reader_version"),
]


@pytest.mark.parametrize("derived,kind", SEMANTIC_CASES)
def test_create_mutate_background_never_runs_recipe(monkeypatch, derived, kind):
    db, files = environment(monkeypatch, derived=derived)
    async def scenario():
        analysis, tasks, original = await create(db, derived=derived)
        mutate(kind, db, files, monkeypatch)
        refreshed = await preflight(db, derived=derived)
        assert not refreshed.executable or refreshed.execution_spec.execution_spec_id != original.execution_spec_id
        if refreshed.executable:
            assert refreshed.execution_spec.authoritative_input_fingerprint != original.authoritative_input_fingerprint
        # A changed authoritative fingerprint must block even the preflight recipe
        # during background revalidation, not just final result publication.
        def forbidden(*args, **kwargs):
            pytest.fail("stale background must not invoke the diagnosis operator")
        monkeypatch.setattr(compiler, "diagnose_conversion", forbidden)
        await tasks()
        assert analysis.status == "failed"
        assert analysis.result_data["status"] in {"stale", "unsupported", "data_invalid"}
        if kind.startswith("base_schema"):
            assert analysis.error_msg == "EXECUTION_SPEC_STALE_OR_MODIFIED"
        if kind == "transform_chain":
            assert analysis.error_msg == "UNSUPPORTED_DERIVED_TRANSFORM_LINEAGE"
    asyncio.run(scenario())


@pytest.mark.parametrize("field", ["display_name", "description", "ui_label", "updated_at", "schema_description", "schema_samples", "join_warnings"])
def test_non_semantic_mutations_keep_spec_and_complete(monkeypatch, field):
    db, _ = environment(monkeypatch, derived=True)
    async def scenario():
        analysis, tasks, original = await create(db, derived=True)
        for ds in db.datasets.values():
            if field == "schema_description":
                ds.schema[0]["description"] = "display only"
            elif field == "schema_samples":
                ds.schema[0]["sample_values"] = ["display preview"]
            elif field == "join_warnings":
                if ds.derivation_plan:
                    ds.derivation_plan["warnings"] = ["display warning changed"]
            else:
                setattr(ds, field, "display metadata changed")
        refreshed = await preflight(db, derived=True)
        assert refreshed.execution_spec == original
        await tasks()
        assert analysis.status == "completed"
    asyncio.run(scenario())


def test_canonical_order_stability(monkeypatch):
    db, _ = environment(monkeypatch, derived=True)
    async def scenario():
        old = (await preflight(db, derived=True)).execution_spec
        for ds in db.datasets.values():
            ds.schema = [dict(reversed(list(c.items()))) for c in reversed(ds.schema)]
            if ds.derivation_plan:
                ds.derivation_plan = dict(reversed(list(ds.derivation_plan.items())))
                ds.derivation_plan["included_dataset_ids"].reverse()
                ds.source_dataset_ids.reverse()
        assert (await preflight(db, derived=True)).execution_spec == old
    asyncio.run(scenario())


VALID = [
    ("boolean", [True, False, True, False]),
    ("numeric", [1, 0, 1, 0]),
    ("string_boolean", ["true", "false", "TRUE", " False "]),
    ("string_binary", ["1", "0", " 1 ", "0"]),
    ("mixed", [True, "false", 1, " 0 "]),
]


@pytest.mark.parametrize("name,values", VALID)
@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_real_reader_explicit_normalization(monkeypatch, extension, name, values):
    content = xlsx_bytes(values) if extension == "xlsx" else csv_bytes(values, quoted=name.startswith("string"))
    if extension == "xlsx":
        book = load_workbook(BytesIO(content))
        for i, value in enumerate(values, start=2):
            cell = book.active.cell(i, 2)
            assert cell.data_type == ("s" if isinstance(value, str) else "b" if isinstance(value, bool) else "n")
            assert cell.value == value
    db, _ = environment(monkeypatch, content=content, extension=extension)
    async def scenario():
        views = await reader.load_capability_dataset_frames(db.datasets["users"])
        raw_values = views.logical_frame.converted.tolist()
        if name in {"string_boolean", "string_binary"}:
            assert all(isinstance(v, str) for v in raw_values)
        analysis, tasks, spec = await create(db)
        record, = spec.authoritative_input_fingerprint.normalization_records
        assert record.policy.version == "1" and record.reader_contract == "capability-logical-reader@1"
        if extension == "csv":
            assert record.string_count == 4 and record.normalization_applied
        else:
            assert record.string_count == sum(isinstance(v, str) for v in values)
            assert record.boolean_count == sum(isinstance(v, bool) for v in values)
            assert record.normalization_applied == any(isinstance(v, str) for v in values)
        await tasks()
        assert analysis.status == "completed"
        assert analysis.result_data["normalization_records"][0] == record.model_dump(mode="json")
        # Other fields use the original reader's behavior.
        legacy = await reader.load_dataset_dataframe(db.datasets["users"])
        assert legacy.drop(columns="converted").equals(views.frame.drop(columns="converted"))
    asyncio.run(scenario())


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
@pytest.mark.parametrize("value", ["yes", "no", "2", "-1", "", "unknown", None, 2, -1])
def test_invalid_raw_value_rejected_without_drop_fill_or_default(monkeypatch, extension, value):
    values = [1, 0, value, 0]
    content = xlsx_bytes(values) if extension == "xlsx" else csv_bytes(values)
    db, _ = environment(monkeypatch, content=content, extension=extension)
    outcome = asyncio.run(preflight(db))
    assert outcome.status == "data_invalid" and not outcome.executable
    assert not db.analyses


@pytest.mark.parametrize("text,native", [("true", True), ("1", 1)])
def test_excel_native_vs_text_changes_fingerprint_even_when_pandas_infers_same(monkeypatch, text, native):
    native_values = [True, False, True, False] if text == "true" else [1, 0, 1, 0]
    text_values = ["true", "false", "true", "false"] if text == "true" else ["1", "0", "1", "0"]
    db, files = environment(monkeypatch, content=xlsx_bytes(native_values), extension="xlsx")
    async def scenario():
        before = (await preflight(db)).execution_spec
        legacy_before = await reader.load_dataset_dataframe(db.datasets["users"])
        files[db.datasets["users"].storage_path] = xlsx_bytes(text_values)
        after = (await preflight(db)).execution_spec
        legacy_after = await reader.load_dataset_dataframe(db.datasets["users"])
        assert legacy_before.converted.tolist() == legacy_after.converted.tolist()
        assert before.execution_spec_id != after.execution_spec_id
        assert before.authoritative_input_fingerprint.normalization_records != after.authoritative_input_fingerprint.normalization_records
    asyncio.run(scenario())


def test_fingerprint_component_inventory_requires_review_on_schema_extension():
    # Changing this list is an explicit reminder to extend the mutation matrix.
    assert set(DatasetSemanticState.model_fields) == {
        "dataset_id", "frame_hash", "schema_semantic_hash", "bound_field_semantics_hash",
        "grain", "row_semantics_hash", "derivation_type", "derivation_fingerprint",
        "transform_chain_fingerprint", "lineage_source_dataset_ids", "base_dataset_id",
        "relationship_snapshot_fingerprint",
    }


def test_changed_bound_column_rejects_queued_spec_before_recipe(monkeypatch):
    content = csv_bytes([0, 1, 1, 0]).decode().splitlines()
    content = [content[0] + ",alternate"] + [line + "," + line.split(",")[1] for line in content[1:]]
    db, _ = environment(monkeypatch, content=("\n".join(content) + "\n").encode())
    async def scenario():
        analysis, tasks, _ = await create(db)
        for binding in analysis.params["request"]["candidate"]["field_bindings"]:
            if binding["role"] == "conversion_flag":
                binding["column"] = "alternate"
        monkeypatch.setattr(compiler, "diagnose_conversion", lambda *a, **k: pytest.fail("modified binding reached recipe"))
        await tasks()
        assert analysis.status == "failed"
        assert analysis.error_msg == "EXECUTION_SPEC_STALE_OR_MODIFIED"
    asyncio.run(scenario())
